from pathlib import Path
import json, textwrap
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
DATA=json.loads((ROOT/'docs/member-demo/guide-content.json').read_text(encoding='utf-8'))
OUT=ROOT/'docs/V-Eats Five Member Command Line Demo Guide.docx'
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=Inches(.65); s.bottom_margin=Inches(.65)
s.left_margin=Inches(.7); s.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(27);doc.styles['Title'].font.bold=True
doc.styles['Subtitle'].font.size=Pt(12)
for name,size in [('Heading 1',19),('Heading 2',13),('Heading 3',11)]:
 doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
 doc.styles[name].paragraph_format.space_before=Pt(10)
 doc.styles[name].paragraph_format.keep_with_next=True
code_style=doc.styles.add_style('SQL Code',1)
code_style.font.name='Consolas';code_style.font.size=Pt(8.6)
code_style.paragraph_format.line_spacing=1.0
code_style.paragraph_format.space_after=Pt(7)
code_style.paragraph_format.keep_together=True
code_style.paragraph_format.left_indent=Inches(.12)

foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('V-Eats command line guide  |  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)

def p(text='',bold=False,style=None):
 para=doc.add_paragraph(style=style);run=para.add_run(text);run.bold=bold
 return para
def h(text,level=2):return doc.add_heading(text,level)
def code(text):
 # Soft-wrap long SQL at whitespace only; identifiers and quoted words stay intact.
 lines=[]
 for line in text.splitlines():
  wrapped=textwrap.wrap(line,width=84,break_long_words=False,break_on_hyphens=False,
      replace_whitespace=False,drop_whitespace=True,subsequent_indent='  ')
  lines.extend(wrapped or [''])
 return p('\n'.join(lines),style='SQL Code')
def page(title):
 heading=h(title,1);heading.paragraph_format.page_break_before=True
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for cell,width in zip(t.rows[0].cells,widths):cell.width=Inches(width)
 for i,text in enumerate(headers):t.rows[0].cells[i].text=text
 for row in rows:
  cells=t.add_row().cells
  for cell,text,width in zip(cells,row,widths):cell.width=Inches(width);cell.text=str(text)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for cell in row.cells:
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcpr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
   tcpr.append(margins)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EBEF');tcpr.append(sh)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.0
    for run in para.runs:run.font.size=Pt(8.5);run.bold=ri==0
 p('')
 return t

p('V Eats command line demonstration guide',style='Title')
p('Individual instructions for Shriya Trisha Tanisi Asmi and Soha',style='Subtitle')
p('Use this guide to perform your Oracle database demonstration in Windows Command Prompt. Every member has an independent section with their own INSERT, UPDATE, DELETE, five SELECT queries, expected results, explanations and code locations. The scripts were executed against the project Oracle database and cleaned up their temporary records.')
h('Who owns each module')
table(['Member','Owned entities'],[(m['name'],m['module']+'\n'+', '.join(m['tables'])) for m in DATA['members']],[1.0,5.8])
h('Read this before rehearsing')
p('1. Complete the shared setup on the next two pages. Then read only your own named section from start to finish. The five sections do not require each other to run first.')
p('2. During the assessed demo, type or paste each SQL statement at SQL> and explain its result before proceeding. Do not demonstrate database queries through the web SQL demo page.')
p('3. The browser is only needed for a brief frontend connectivity check if the examiner requests it. Oracle queries and DML in this guide run directly in SQL*Plus.')
h('What to say about JDBC')
p('Say: "Our Next.js backend uses the node-oracledb Thin driver, which is the Oracle connectivity mechanism for Node.js. We do not have a Java JDBC file. The shared connection pool is in lib/oracle.ts, and the shared CRUD SQL is in app/api/manage/route.ts."')
p('The assignment allows JDBC, ODBC or an equivalent mechanism. Do not claim this project uses Java JDBC. The five command-line SELECTs are presentation queries; they are not all automatically executed by the frontend.')
p('Reference snapshot: '+DATA['sourceCommit']+' plus the accompanying member worksheets. File and line references were checked against the working code. Later edits can shift line numbers.')

page('Shared setup on a fresh laptop')
p('If your laptop already runs V-Eats with a working .env.local and Oracle container, skip to Shared command prompt instructions. Otherwise install Git, Node.js 24 and Docker Desktop, start Docker using Linux containers, then open Windows Command Prompt by searching for cmd in Start.')
h('Clone into your Desktop folder')
code('cd /d "%USERPROFILE%\\Desktop"\ngit clone --branch dev https://github.com/shriyapsawant21-wq/V-Eats.git\ncd V-Eats\nnpm ci')
p('If the Desktop folder is redirected to OneDrive, use File Explorer to copy your real folder path, then type cd /d "that path". If V-Eats is already cloned, do not clone over it; use git pull --ff-only origin dev inside that folder after keeping local edits safe.')
h('Create Oracle once')
code('docker run -d --name oracle-free --restart unless-stopped ^\n -p 127.0.0.1:1521:1521 -e ORACLE_RANDOM_PASSWORD=yes ^\n -v veats-oracle-data:/opt/oracle/oradata ^\n --health-cmd healthcheck.sh --health-interval 10s ^\n --health-timeout 5s --health-retries 30 ^\n --health-start-period 120s gvenzl/oracle-free:23.26.2-slim')
p('Copy that entire block into Windows Command Prompt. The ^ characters continue the command on the next line; do not paste this CMD block into PowerShell. If a compatible oracle-free container already exists, use docker start oracle-free instead. Do not delete an existing database volume.')
code('docker exec oracle-free healthcheck.sh')
p('Wait for that command to succeed; the first database startup can take several minutes. Then create the isolated application schema and sample data:')
code('npm run db:setup\nnpm run seed\nnpm run db:audit')
p('Setup writes your own ORACLE_PASSWORD and ADMIN_TOKEN into .env.local. Keep this file private. Do not replace its real values with placeholders from .env.example. Each laptop has its own credentials and independent database data.')
h('Use the accompanying SQL worksheets')
p('Extract the supplied SQL bundle into the V-Eats repository so sql/members/shriya.sql and the other four named files exist. The main guide works without those files because every SQL statement is printed here. Optional paused worksheets require them.')

page('Shared command prompt instructions')
h('Know which prompt you are using')
table(['Prompt','What goes there'],[('C:\\...>','Windows commands such as cd, docker, npm and code.'),('SQL>','Oracle SQL, COMMIT, SAVEPOINT, DESCRIBE and EXIT.'),('2   3   4','SQL*Plus continuation lines while you enter a multi-line SQL statement. These are normal.')],[1.0,5.8])
p('Never type the displayed C:\\...> or SQL> prompt itself. SQL statements end with a semicolon. Strings use straight single quotes. Copy plain SQL, not the explanations or the expected output.')
h('Start a backend SQL session')
code('cd /d "%USERPROFILE%\\Desktop\\V-Eats"\ndocker start oracle-free\ndocker exec oracle-free healthcheck.sh\ndocker exec -it oracle-free sqlplus /nolog')
p('The final command opens SQL*Plus. At SQL>, type:')
code('CONNECT VEATS_APP@//127.0.0.1:1521/FREEPDB1')
p('At Enter password:, privately paste ORACLE_PASSWORD from .env.local and press Enter. Nothing is shown while you type or paste the password; that is normal. Do not use ADMIN_TOKEN. Once Connected appears, run:')
code('SHOW USER\nSET LINESIZE 180\nSET PAGESIZE 100')
p('SHOW USER must report VEATS_APP. Stop if it says SYS or SYSTEM. The backend demo does not need npm run dev; only Oracle must be running.')
h('What to do when something goes wrong')
p('ORA-00001 means the ID/name already exists, often after an interrupted rehearsal. Check the exact named demo ID, and use only your own cleanup statements. ORA-02291 means a referenced parent ID is missing: check the seeded IDs. ORA-02292 means child rows still reference a parent: delete links first. ORA-01017 means the Oracle password does not match this database. Connection refused means Docker/Oracle is not ready.')
p('An error is not a successful result. Stop, read it and correct it. ROLLBACK only undoes uncommitted work. Your earlier COMMITs survive EXIT and container restarts. Do not repeatedly paste a failed INSERT.')
h('Show code without guessing line numbers')
p('In VS Code press Ctrl+P, type the exact relative file path from your section and press Enter. Press Ctrl+G, type the indicated line number and press Enter. The right side of the screen should show the source; keep Command Prompt on the left. Never project .env.local.')

def verified(m,step):
 out=next(x for x in DATA['results']['members'] if x['name']==m['name'])
 match=next(x for x in out['outputs'] if x['title']==step['title'])
 return match

def query_block(m,step):
 h(step['title'],2)
 p(f"SQL file: {step['file']}  |  Lines {step['line']} to {step['endLine']}")
 code(step['sql']+';')
 p('Explain: '+step['why'])
 result=verified(m,step)['results'][0]
 p(f"Verified result: {len(result['rows'])} row(s) on the seeded test database. "+step['expected'])

for idx,m in enumerate(DATA['members'],1):
 page(f"Member {idx} {m['name']} {m['module']}")
 p('This section belongs to '+m['name']+'. Do not execute another member\'s fixture IDs. Replace the example repository path only if your clone is in a different folder.')
 h('Your responsibility')
 p(m['introduction']);p(m['columns']);p('Before you start: '+m['prerequisites'])
 h('Open your own command line session')
 code('cd /d "%USERPROFILE%\\Desktop\\V-Eats"\ndocker start oracle-free\ndocker exec oracle-free healthcheck.sh\ndocker exec -it oracle-free sqlplus /nolog')
 p('At SQL>, connect using your local Oracle password when prompted:')
 formats={'shriya':[('restaurant_id','A12'),('restaurant_name','A28'),('block','A10'),('landmark','A28'),('outlet','A28'),('location','A28')],
 'trisha':[('food_id','A12'),('food_name','A25'),('description','A43'),('category_name','A20'),('cuisine_name','A20'),('type_name','A20')],
 'tanisi':[('kind','A10'),('id','A15'),('name','A27'),('category_name','A27'),('cuisine_name','A27'),('type_name','A25'),('food_name','A25')],
 'asmi':[('food_name','A25'),('allergen_name','A26'),('listed_allergens','A40')],
 'soha':[('restaurant_name','A28'),('food_name','A25'),('formatted_price','A16')]}[m['slug']]
 code('CONNECT VEATS_APP@//127.0.0.1:1521/FREEPDB1\nSHOW USER\nSET LINESIZE 180\nSET PAGESIZE 100\n'+'\n'.join('COLUMN '+name+' FORMAT '+fmt for name,fmt in formats))
 p('Verify VEATS_APP is the connected user. You do not need the browser admin key for this session. The password is ORACLE_PASSWORD in .env.local; keep it off the projector.')
 h('Check for an interrupted rehearsal')
 table_name='food_item' if m['slug']=='trisha' else 'allergen' if m['slug']=='asmi' else 'category' if m['slug']=='tanisi' else 'restaurant'
 key='food_id' if m['slug']=='trisha' else 'allergen_id' if m['slug']=='asmi' else 'category_id' if m['slug']=='tanisi' else 'restaurant_id'
 code(f"SELECT COUNT(*) AS already_exists FROM {table_name}\nWHERE {key}='{m['ids'][0]}';")
 p('Expected: 0. If it is 1, do not overwrite the row. Finish your own cleanup at the end of this section, confirm zero, then restart your module.'+(' Also check TANISI_U in CUISINE and TANISI_D in DIETARY_TYPE if a previous run stopped partway.' if m['slug']=='tanisi' else ''))
 h('What to say at your introduction')
 p('"'+m['introduction']+'"')

 page(m['name']+' insert and update')
 p('Enter these blocks in order at SQL>. When a block contains several statements, each semicolon executes the statement before it. Explain the affected-row message after every statement.')
 for step in m['writes']:
  h(step['title']);code(step['sql']+';');p('Expected: '+step['expected']+' '+step['why'])
 h('After these writes')
 p('Run S1 on the next page to prove the row exists and the update took effect. Your original seed records remain unchanged. Do not run DELETE yet; your SELECTs should run while your demonstration record exists.')
 p('COMMIT makes your change visible to another session. If a frontend connectivity check is requested, refresh the actual app after COMMIT.'+(' A food item without a Contains row does not appear as an outlet offer.' if m['slug']=='trisha' else ''))

 page(m['name']+' SELECT queries one to three')
 for step in m['selects'][:3]:query_block(m,step)

 page(m['name']+' SELECT queries four and five')
 for step in m['selects'][3:]:query_block(m,step)
 h('Questions you must be able to answer')
 for question,answer in m['viva']:
  p(question,bold=True);p(answer)

 if m.get('transactions'):
  page(m['name']+' transactions')
  p('Do this after S1 to S5 and before cleanup. Price 95 was committed earlier. Do not issue COMMIT or any DDL between SAVEPOINT and ROLLBACK TO.')
  for step in m['transactions']:
   h(step['title']);code(step['sql']+';');p('Expected: '+step['expected'])
  h('What to explain before continuing')
  p('The original 95 price was committed. The 999 update was not committed. ROLLBACK TO restores the value at the savepoint but does not remove the row inserted in the earlier committed transaction. A COMMIT here would erase savepoints, so the order matters.')
  p('Next perform your cleanup on the following page. Your separate scratch DDL page comes after cleanup; continue forward through the document.')

 page(m['name']+' cleanup and source code')
 h('Delete only your own demonstration records')
 for step in m['deletes']:
  code(step['sql']+';');p('Expected: '+step['expected']+' '+step['why'])
 p('After the remaining count is zero, '+('stay in SQL*Plus and continue to your scratch DDL page after the source references below.' if m['slug']=='soha' else 'type EXIT; to leave SQL*Plus. The next member can now start independently.'))
 h('Where your connectivity and SQL code lives')
 p('There is no Java JDBC file per person. The backend shares one Oracle connection layer and one generic CRUD handler; your entity definition selects the correct table and key. '+m['reportExplanation'])
 compact=[]
 schema_refs=[r for r in m['refs'] if r['file']=='sql/01-schema.sql']
 compact.append(('sql/01-schema.sql\nLines '+', '.join(r['lines'] for r in schema_refs),'Your table definitions and database constraints: '+', '.join(m['tables'])))
 for r in m['refs']:
  if r['file']=='sql/01-schema.sql':continue
  if r['file']=='lib/oracle.ts':
   if not any(x[0].startswith('lib/oracle.ts') for x in compact):compact.append(('lib/oracle.ts\nLines 1, 5 and 9','Oracle driver import, connection pool and SELECT execution with binding and connection release'))
  elif r['file']=='app/api/manage/route.ts' and ('builder' in r['what'] or 'Relationship UPDATE' in r['what']):
   if not any(x[0].startswith('app/api/manage/route.ts\nSQL') for x in compact):
    inserts=next(x['lines'] for x in m['refs'] if 'INSERT builder' in x['what'])
    updates=next(x['lines'] for x in m['refs'] if 'UPDATE' in x['what'])
    deletes=next(x['lines'] for x in m['refs'] if 'DELETE builder' in x['what'])
    compact.append((f'app/api/manage/route.ts\nSQL lines {inserts}, {updates}, {deletes}',f'INSERT line {inserts}; UPDATE line {updates}; DELETE line {deletes}. Your entity selects the table and key.'))
  else:compact.append((r['file']+'\nLines '+r['lines'],r['what']))
 table(['File and verified lines','What to show'],compact,[2.8,4.0])
 p('When asked where your INSERT is: show your lib/entities.ts entry, then the shared INSERT line in app/api/manage/route.ts. Point to execute, commit and rollback next. The literal SQL you typed for this assessment is in your named sql/members worksheet.')
 p('A colon placeholder such as :restaurant_id is a bind variable. The driver receives values separately from SQL; it is the Node.js equivalent of prepared parameter binding, not string-concatenated user input.')
 if m['slug']=='soha':
  page('Soha scratch DDL and exit')
  p('Your menu and parent fixture should now be deleted and committed. Use only this dedicated scratch table. It is disposable and is not one of the eight project entities. Oracle DDL implicitly commits; do not use project tables for this demonstration.')
  for step in DATA['ddlSteps']:
   h(step['title']);code(step['sql']+';');p('Expected: '+step['expected']+' '+step['why'])
  h('Finish your session')
  code('EXIT;')
  p('SQL*Plus closes and returns to the Windows command prompt. Your menu fixture and scratch table have been removed.')

page('Final rehearsal checklist')
p('Every member should independently repeat their module until they can explain each query without reading the explanation. Do not memorize only the command that runs an entire script.')
for text in ['I can open Command Prompt and connect as VEATS_APP without showing my password.','I can identify my owned tables, primary keys, foreign keys and relevant ER relationships.','I can execute my INSERT, UPDATE and DELETE and interpret the affected-row counts.','I can execute all five SELECTs and explain joins, nesting, grouping, functions or set operations used in them.','I know my literal worksheet path, entity definition, shared CRUD lines, pool line and execute/commit/rollback lines.','I can recover from an interrupted rehearsal using only my own cleanup IDs.','I leave the remaining count at zero and do not alter or delete seed data.']:
 p('[ ] '+text)
h('Optional paused worksheets')
p('Manual statement-by-statement execution is the primary assessment method. If you want a paused rehearsal script, extract the SQL bundle into V-Eats and run these Windows commands before opening SQL*Plus:')
code('docker exec oracle-free mkdir -p /tmp/veats-members\ndocker cp sql/members/. oracle-free:/tmp/veats-members')
p('Then connect to SQL*Plus and run only your own file, for example:')
code('@/tmp/veats-members/shriya.sql')
p('Replace shriya with trisha, tanisi, asmi or soha. The worksheet pauses after every step so you can explain the output and press Enter. Do not use an all-members runner in place of understanding your statements.')
h('Keep the claims accurate')
p('The application uses Oracle through node-oracledb Thin, not Java JDBC. The backend demo is in SQL*Plus. Original schema qualifications include the corrected Contains Restaurant_ID, the many-to-many allergen mapping, JSON MUST_HAVE storage, supporting application fields and unenforced parent-side minimum participation. The actual Laboratory Cycle Sheets I to VI still need to be checked against your college copies.')
p('Sample menu prices and allergen links are academic demonstration data. Do not describe them as current verified campus prices or medical allergy advice. Assign names to genuine responsibilities during your presentation; do not claim that generated examples prove individual authorship.')

doc.core_properties.title='V Eats command line demonstration guide'
doc.core_properties.subject='Five independent Oracle modules and source code references'
doc.core_properties.author='V-Eats project team'
for root in [doc._element,doc.styles.element]:
 for border in list(root.xpath('.//w:pBdr')):border.getparent().remove(border)
doc.save(OUT)
print(OUT)
